#!/usr/bin/env python3
"""Suggest papers and figures from a Zotero library for an article draft (research
or review), section by section, with a match %, a note on whether the draft
already cites each paper, corresponding-author emails and draft permission emails.
These are suggestions: nothing is removed, weaker matches are only listed lower.

Usage:  python3 zotero_filter.py review.docx library.csv [--figs 3]
Export: Zotero > File > Export Library > Format: CSV  (keeps the PDF paths)
        or skip the export: give "zotero" or "zotero:Collection name" instead of
        library.csv to read straight from a running Zotero 7+ (Settings > Advanced >
        "Allow other applications on this computer to communicate with Zotero")
Draft:  .docx, .md, .tex or .txt (split on headings; no headings = one section)
Needs:  Poppler (pdftotext, pdftoppm) for figures and emails, either on PATH or in a
        "poppler" folder next to this file; internet for OpenAlex
Runs on macOS, Windows and Linux.
Output: <library>_report/report.html and papers.csv

Before judging a paper the tool reads its abstract and its introduction (taken
from the PDF), not the abstract alone. Match % is similarity between a section
of the draft and that text, not a probability of relevance. With word matching, 100% is the best-matching
paper in the library. strong match >= 60%, possible match >= 40%, weak match
below that (no abstract and < 40% = check by hand).

Optional local model: --semantic adds SPECTER2, a model trained on scientific
papers that runs on this computer (no key, nothing sent out). It scores by
meaning, so synonyms count, and is averaged with the word-match score for
papers. Figures stay on word matching, which did as well on captions in testing.
Needs:  pip install adapters   (and ~840 MB of model files on first use)

Optional AI mode: point the tool at any OpenAI-style chat endpoint by giving
LLM_BASE_URL and LLM_MODEL (plus LLM_API_KEY unless the model is local) in any
of three ways, highest priority first:
    1. options:   --llm-url URL --llm-model NAME [--llm-key KEY]
    2. terminal:  export LLM_BASE_URL=... LLM_MODEL=... LLM_API_KEY=...
    3. file:      llm.env next to this script or in the CiteFilter settings folder
                  (or --llm-file PATH), one NAME=value per line; copy
                  llm.env.example to start
An AI model then scores each paper 0-100 (same 60/40 cut-offs), picks the
figures and personalises the emails. Excerpts of the draft, the abstracts and
the captions are sent to that endpoint. Unset, or on any AI failure, the tool
falls back to word matching and the fixed email template.
"""
import argparse
import csv
import functools
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

VERSION = "1.0.0"
APP_DIR = Path(__file__).parent  # also where a packaged app keeps its bundled files
GOOD, MEDIUM = 0.6, 0.4  # match % from which a paper is a strong / possible match
LABELS = {"strong": "strong match", "possible": "possible match", "weak": "weak match",
          "check": "check by hand"}
REFERENCES = re.compile(r"(?i)^\W*(?:\d+\W+)?(references?|bibliography|works cited|literature cited|reference list)\W*$")
# Local model. SPECTER2 gives even unrelated papers a cosine near 0.85, so scores
# are stretched from that floor to the best match. ponytail: floor calibrated on a
# 12-paper test library; lower it if clearly relevant papers show 0% meaning match.
SEMANTIC_FLOOR = 0.85
CHUNK_WORDS, CHUNK_STEP = 300, 250  # the model reads ~512 tokens, so long sections go in pieces
# What is read from each paper's PDF before judging it
INTRO_WORDS = 900  # how much of the introduction is read
INTRO_HEAD = re.compile(r"(?im)^[ \t]*(?:[1I][.)]?[ \t]+)?introduction[ \t]*$")
INTRO_END = re.compile(r"(?m)^[ \t]*(?:(?:2|II)[.)]?[ \t]+[A-Z][^\n]{2,60}|(?:\d[.)]?[ \t]+)?(?:Experimental(?: Section)?"
                       r"|Materials and Methods|Methods|Methodology|Theory|Results(?: and Discussion)?|Background"
                       r"|Related Work)[ \t]*)$")
MIN_WORDS = 30  # sections shorter than this match on noise, skip them
OPENING_WORDS = 100  # same for text before the first heading
DRAFT_TYPES = (".docx", ".md", ".tex", ".txt")
SECTION_NAMES = ("abstract", "introduction", "background", "related work", "theory", "methods", "methodology",
                 "materials and methods", "experimental", "experimental section", "results",
                 "results and discussion", "discussion", "conclusion", "conclusions", "summary", "outlook",
                 "acknowledgements", "acknowledgments", "references", "bibliography")
HEADING = re.compile(r"^(?:#+\s+|\\(?:sub)*section\*?\{)(.+?)\}?\s*$")
PAPER_FIELDS = ("Title", "Abstract Note", "Manual Tags", "Automatic Tags")
# Captions: "Fig. 3." / "Figure 3:" / "Fig. 3 |" / "FIG. 3." / "Fig. 3 Capitalised"
FIG = r"((?:Fig(?:ure)?|FIG(?:URE)?)\.?\s*(S?\d+)(?:\s*[.:|]\s*|\s+(?=[A-Z(])))"
# A caption that starts its own text block is trusted over one that only starts
# a line, because "Fig. 3. The ..." can also open a line of running text.
CAPTION_PATTERNS = (re.compile(r"(?:\A|\n[ \t]*\n)[ \t]*" + FIG), re.compile("^" + FIG, re.M))
# A figure the paper itself borrowed: permission must come from the original source.
BORROWED = re.compile(r"(?i)reproduced|reprinted|adapted (?:from|with)|with permission|copyright|©")

BLOCK = re.compile(r'<block xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</block>', re.S)
BODY_WORDS = 12  # a text block this long is running text, not a label inside a figure

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
NOT_AUTHOR = ("permission", "reprint", "journal", "support", "info", "editor", "office", "help", "contact")
CORRESPONDING = re.compile(r"(?i)correspond|e-?mail\s*:")
FREE_LICENSES = ("cc-by", "cc0", "public-domain")

ZOTERO_API = "http://localhost:23119/api/users/0"  # Zotero's local API; 0 = the local user
ZOTERO_HELP = ('Open Zotero (version 7 or later), then in Settings > Advanced tick "Allow other '
               'applications on this computer to communicate with Zotero".')
ZOTERO_COLUMNS = ("Title", "Author", "Publication Title", "Publication Year", "DOI", "Abstract Note",
                  "Manual Tags", "Automatic Tags", "File Attachments")

MAIL_BODY = """Dear {name},

I am [YOUR NAME] from [YOUR AFFILIATION]. I am preparing an article titled
"[ARTICLE TITLE]" for submission to [TARGET JOURNAL].

I would like to request your permission to reproduce the following figure:

    Figure {fig} from: {citation}

The figure would appear in the print and electronic versions of the article, with
full citation and the credit line "Reproduced with permission from [REF]".

If the copyright is held by the publisher rather than by you, I would be grateful
if you could let me know so that I can direct the request there.

Thank you for your time and for your work.

Kind regards,
[YOUR NAME]
[YOUR AFFILIATION]
[YOUR EMAIL]"""

# AI mode. ponytail: only the start of each section / abstract is sent, enough
# to convey the topic and small enough for local models; raise if sections are
# long and change subject midway.
SECTION_WORDS, ABSTRACT_WORDS = 250, 300
BATCH = 10  # papers rated per AI call (each carries its abstract and the start of its introduction)
INTRO_AI_WORDS = 200  # how much of each introduction is sent to the AI
CANDIDATES = 10  # word-matched captions per section offered to the AI
RATE_SYSTEM = """You help a researcher writing a research or review article decide which papers in their \
reference library are relevant to it. You get the start of each section of the draft and a list \
of papers, each with its title, abstract and the start of its introduction. Read both before \
judging. Score every paper from 0 to 100 for how useful it is to cite in this article: \
80-100 directly about a section's topic, 60-79 clearly relevant, 40-59 related background, \
below 40 not relevant. Judge by meaning, not shared words: synonyms and abbreviations count. \
Reply with only a JSON array, one object per paper: \
[{"id": 0, "score": 85, "section": "<heading of the best-fitting section, copied exactly>", \
"reason": "<one sentence, at most 20 words>"}]"""
FIG_SYSTEM = """You help a researcher choose figures to reproduce in their article. You get \
one section of the draft and candidate figure captions from papers in their library. Pick the \
figures that would best illustrate this section for a reader, best first. Prefer schematics and \
summary figures over raw data plots when both fit. Pick fewer than asked if the rest do not fit. \
Reply with only a JSON array: [{"id": 3, "reason": "<one sentence, at most 20 words>"}]"""
MAIL_SYSTEM = """You draft a short, polite email from a researcher to the corresponding \
author of a paper, asking permission to reproduce one figure in an article they are writing. Write only \
the email body, starting with the greeting. Under 200 words, plain text. Say specifically why \
this figure suits the section, using only the facts given: do not invent results, \
praise or details about the paper. State that the figure would appear in the print and \
electronic versions with full citation and a credit line. Ask the author to say so if the \
publisher holds the copyright. The researcher's details are not known yet, so keep these \
placeholders exactly as written wherever they belong: [YOUR NAME], [YOUR AFFILIATION], \
[ARTICLE TITLE], [TARGET JOURNAL], [YOUR EMAIL]. Sign off with [YOUR NAME], [YOUR AFFILIATION] \
and [YOUR EMAIL] on separate lines."""


def config_dir():
    """Folder for CiteFilter's own settings (llm.env), in the place each system expects."""
    if os.name == "nt":
        base = Path(os.environ.get("APPDATA") or Path.home())
    elif sys.platform == "darwin":
        base = Path.home() / "Library" / "Application Support"
    else:
        base = Path(os.environ.get("XDG_CONFIG_HOME") or Path.home() / ".config")
    return base / "CiteFilter"


def output_dir():
    """Folder for libraries loaded from Zotero and their reports."""
    documents = Path.home() / "Documents"
    # Windows users look in Documents; on macOS that folder triggers a permission prompt, so use home
    return (documents if os.name == "nt" and documents.is_dir() else Path.home()) / "CiteFilter"


@functools.lru_cache(maxsize=None)
def tool(name):
    """Path of a Poppler program: the copy shipped in ./poppler if there is one, else the one on PATH."""
    exe = name + (".exe" if os.name == "nt" else "")
    shipped = sorted((APP_DIR / "poppler").rglob(exe)) if (APP_DIR / "poppler").is_dir() else []
    return str(shipped[0]) if shipped else shutil.which(name)


def run_tool(args):
    """Run a Poppler program quietly and return its output as text."""
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0  # no console window flashing up on Windows
    return subprocess.run([str(a) for a in args], capture_output=True, text=True, encoding="utf-8",
                          errors="ignore", creationflags=flags, check=False).stdout


def looks_like_heading(text, bold=False):
    """Guess a heading in a draft that uses no heading styles: a short line that is all bold,
    numbered like "2. Methods", or a standard section name."""
    words = text.split()
    if not words or len(words) > 12 or text.rstrip().endswith((".", ",", ";")):
        return False
    return bool(bold or re.match(r"\d+(\.\d+)*[.)]?\s+[A-Z]", text.strip())
                or text.strip().rstrip(":").lower() in SECTION_NAMES)


def is_subsection(text):
    """A line like "2.3 Attenuation mechanisms": a heading even when the author skipped the heading style."""
    return bool(len(text.split()) <= 12 and not text.rstrip().endswith((".", ",", ";"))
                and re.match(r"\d+\.\d+(\.\d+)*\.?\s+[A-Z]", text.strip()))


def read_lines(path):
    """Return [(is_heading, text)] for a draft file."""
    path = Path(path)
    if not path.is_file():
        raise SystemExit(f"Draft not found: {path}")
    if path.suffix.lower() not in DRAFT_TYPES:
        raise SystemExit(f"The draft must be a .docx, .md, .tex or .txt file, not \"{path.name}\". "
                         "In Word, use File > Save As and choose Word Document (.docx).")
    if path.suffix.lower() == ".docx":
        import docx

        try:
            paragraphs = docx.Document(str(path)).paragraphs
        except Exception as err:  # not really a Word file, or damaged
            raise SystemExit(f"Could not open {path.name} as a Word document ({type(err).__name__}). "
                             "Open it in Word and save it again as .docx.") from err
        styled = [p.style.name.startswith(("Heading", "Title")) for p in paragraphs]
        lines = [(is_styled or is_subsection(p.text), p.text) for is_styled, p in zip(styled, paragraphs)]
        if not any(styled):  # headings typed by hand, not with styles
            lines = [(looks_like_heading(p.text, bold=all(r.bold for r in p.runs if r.text.strip())
                                         and any(r.text.strip() for r in p.runs)), p.text) for p in paragraphs]
        return lines
    lines, marked = [], False
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        m = HEADING.match(line.strip())
        marked = marked or bool(m)
        lines.append((bool(m) or is_subsection(line), m.group(1) if m else line))
    if not marked:  # no # or \\section markers: guess headings from how lines look
        lines = [(looks_like_heading(text), text) for _, text in lines]
    return lines


def split_sections(lines):
    """Return [(heading, body)]; falls back to one section if no usable headings."""
    whole = [("Whole article", " ".join(text for _, text in lines))]
    if not any(is_heading for is_heading, _ in lines):
        return whole
    sections, title, body = [], None, []
    for is_heading, text in lines + [(True, "")]:
        if not is_heading:
            body.append(text)
            continue
        joined = " ".join(body)
        # text before the first heading is mostly title, authors and addresses: keep it only if long
        if len(joined.split()) >= (MIN_WORDS if title else OPENING_WORDS):
            sections.append((title or "Opening text (before the first heading)", joined))
        title, body = text.strip() or title, []
    return sections or whole


def split_references(lines):
    """Split draft lines into (lines without the reference list, reference-list entries)."""
    best = (lines, [])
    for i, (_, text) in enumerate(lines):
        if REFERENCES.match(text):  # a draft can hold an empty "References" stub too: keep the fullest
            rest = lines[i + 1:]
            end = next((k for k, (is_heading, _) in enumerate(rest) if is_heading), len(rest))
            entries = [t.strip() for _, t in rest[:end] if t.strip()]
            if len(entries) > len(best[1]):
                best = (lines[:i] + rest[end:], entries)
    return best


def norm(text):
    return " ".join(re.findall(r"[a-z0-9]+", text.lower()))


def cited_numbers(text):
    """Numbers cited in square brackets, e.g. "[3, 7-9]" -> {3, 7, 8, 9}."""
    numbers = set()
    for group in re.findall(r"\[([\d,;\s\u2013-]+)\]", text):
        for part in re.split(r"[,;]", group.replace("\u2013", "-")):
            low, _, high = part.strip().partition("-")
            if low.strip().isdigit():
                top = int(high) if high.strip().isdigit() else int(low)
                numbers |= set(range(int(low), min(top, int(low) + 200) + 1))
    return numbers


def find_citations(sections, references, papers):
    """Note on each paper where the draft already cites it.

    Sets p["cited_in"] = {section index: "sure" | "likely"} and p["in_refs"] = "sure" | "likely" | "".
    "sure" = DOI or full title found in the reference list; "likely" = only first author and year match.
    """
    # ponytail: text matching. Misses superscript-number citations with no reference list, and
    # "likely" can confuse two papers by the same first author in one year. A reference manager's
    # own citation data would be exact; this needs none.
    numbered = [(int(m.group(1)) if (m := re.match(r"\W*(\d{1,3})[\].)]", entry)) else k, entry)
                for k, entry in enumerate(references, 1)]
    numbers_in = [cited_numbers(body) for _, body in sections]
    for p in papers:
        title, doi = norm(p["Title"]), (p.get("DOI") or "").lower()
        author = (p.get("Author") or "").split(";")[0].strip()
        surname = (author.split(",")[0] if "," in author else author.split()[-1] if author else "").strip()
        year = (p.get("Publication Year") or "").strip()
        by_author = (re.compile(rf"\b{re.escape(surname)}\b(?:\s+et\s+al\.?|,?\s+(?:and|&)\s+[A-Z][\w-]+|(?:,\s+[A-Z][\w-]+)+)?"
                                rf"[,\s]*\(?\s*{re.escape(year)}[a-z]?\b")
                     if len(surname) >= 3 and year else None)
        refs = {}  # reference number -> how sure
        for number, entry in numbered:
            if (doi and doi in entry.lower()) or (len(title) >= 20 and title in norm(entry)):
                refs[number] = "sure"
            elif by_author and year in entry and re.search(rf"\b{re.escape(surname)}\b", entry):
                refs.setdefault(number, "likely")
        p["in_refs"] = "sure" if "sure" in refs.values() else "likely" if refs else ""
        p["cited_in"] = {}
        for j, (_, body) in enumerate(sections):
            hits = [how for number, how in refs.items() if number in numbers_in[j]]
            if "sure" in hits:
                p["cited_in"][j] = "sure"
            elif hits or (by_author and by_author.search(body)):
                p["cited_in"][j] = "likely"


def read_papers(path):
    if not Path(path).is_file():
        raise SystemExit(f"Library file not found: {path}")
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        return [r for r in csv.DictReader(f) if r.get("Title")]


class ZoteroError(Exception):
    """Zotero is not running, its local API is off, or it refused a request."""


def zotero_request(path):
    """GET one Zotero local API URL; return (body bytes, headers)."""
    request = urllib.request.Request(ZOTERO_API + path, headers={"Zotero-API-Version": "3"})
    try:
        with urllib.request.urlopen(request, timeout=20) as response:
            return response.read(), response.headers
    except urllib.error.HTTPError as err:
        detail = err.read(200).decode(errors="ignore").strip()
        raise ZoteroError(f"Zotero refused the request (HTTP {err.code}: {detail}). {ZOTERO_HELP}") from err
    except OSError as err:
        raise ZoteroError(f"Could not reach Zotero. {ZOTERO_HELP}") from err


def zotero_list(path):
    """Fetch a whole list. The local API sends everything at once, but page anyway if it does not."""
    results = []
    while True:
        body, headers = zotero_request(f"{path}{'&' if '?' in path else '?'}start={len(results)}")
        try:
            page = json.loads(body)
        except ValueError as err:
            raise ZoteroError("Zotero sent a reply that could not be read.") from err
        results += page
        if not page or len(results) >= int(headers.get("Total-Results") or 0):
            return results


def zotero_collections():
    """Return [(key, "Parent / Child")] for every collection, sorted by name."""
    data = {c["key"]: c["data"] for c in zotero_list("/collections")}

    def full_name(key):
        parent = data[key].get("parentCollection")
        return (full_name(parent) + " / " if parent in data else "") + data[key]["name"]

    return sorted(((key, full_name(key)) for key in data), key=lambda c: c[1].lower())


def zotero_pdf_path(attachment):
    """Return the path of a PDF attachment on this computer, or None."""
    if attachment["data"].get("contentType") != "application/pdf":
        return None
    href = ((attachment.get("links") or {}).get("enclosure") or {}).get("href", "")
    if not href.startswith("file:"):
        try:
            href = zotero_request(f"/items/{attachment['key']}/file/url")[0].decode().strip()
        except ZoteroError:
            return None  # e.g. the file was never downloaded to this computer
    path = urllib.request.url2pathname(urllib.parse.urlparse(href).path)  # also turns /C:/... into C:\...
    return path if Path(path).is_file() else None


def zotero_row(item, pdf_paths):
    """Shape one Zotero item like a row of Zotero's own CSV export."""
    data = item["data"]
    authors = [c.get("name") or f"{c.get('lastName', '')}, {c.get('firstName', '')}".strip(", ")
               for c in data.get("creators") or [] if c.get("creatorType", "author") == "author"]
    year = re.search(r"\d{4}", data.get("date") or "")
    return {"Title": data.get("title") or "", "Author": "; ".join(authors),
            "Publication Title": data.get("publicationTitle") or data.get("bookTitle") or "",
            "Publication Year": year.group(0) if year else "", "DOI": data.get("DOI") or "",
            "Abstract Note": data.get("abstractNote") or "",
            "Manual Tags": "; ".join(t.get("tag", "") for t in data.get("tags") or []),
            "Automatic Tags": "", "File Attachments": "; ".join(pdf_paths)}


def zotero_export(collection=None, name="Whole library"):
    """Save the library (or one collection) as a Zotero-style CSV; return (path, papers, with PDF)."""
    top = zotero_list((f"/collections/{collection}" if collection else "") + "/items/top")
    papers = {i["key"]: i for i in top if i["data"].get("itemType") not in ("attachment", "note", "annotation")}
    pdfs = {key: [] for key in papers}
    for attachment in zotero_list("/items?itemType=attachment"):
        parent = attachment["data"].get("parentItem")
        if parent in pdfs and (path := zotero_pdf_path(attachment)):
            pdfs[parent].append(path)
    folder = output_dir()
    folder.mkdir(parents=True, exist_ok=True)
    out = folder / f"zotero_{re.sub(r'[^A-Za-z0-9]+', '_', name).strip('_') or 'library'}.csv"
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=ZOTERO_COLUMNS)
        writer.writeheader()
        writer.writerows(zotero_row(papers[key], pdfs[key]) for key in papers)
    return out, len(papers), sum(1 for key in papers if pdfs[key])


def zotero_library(spec):
    """Turn "zotero" or "zotero:Collection name" into the path of a freshly saved CSV."""
    wanted = spec.partition(":")[2].strip().lower()
    key, name = None, "Whole library"
    if wanted:
        collections = zotero_collections()
        match = [c for c in collections if wanted in (c[1].lower(), c[1].lower().rsplit(" / ", 1)[-1])]
        if not match:
            raise ZoteroError(f"No Zotero collection named {wanted!r}. Found: "
                              + (", ".join(n for _, n in collections) or "none"))
        key, name = match[0]
    path, count, with_pdf = zotero_export(key, name)
    print(f"Loaded {count} papers from Zotero ({name}); {with_pdf} have a PDF on this computer.")
    return path


def similarity(queries, docs):
    """TF-IDF cosine similarity, shape (len(docs), len(queries))."""
    # ponytail: TF-IDF matches shared words, not meaning. Swap for embeddings
    # or an LLM rerank if relevant items with different vocabulary rank low.
    tfidf = TfidfVectorizer(stop_words="english", sublinear_tf=True,
                            ngram_range=(1, 2)).fit_transform(queries + docs)
    return cosine_similarity(tfidf[len(queries):], tfidf[:len(queries)])


def load_specter():
    """Return embed([(title, text)]) -> unit vectors using SPECTER2, or None if unavailable."""
    os.environ.setdefault("HF_HUB_DISABLE_PROGRESS_BARS", "1")
    try:
        import torch
        import transformers
        from adapters import AutoAdapterModel

        transformers.logging.set_verbosity_error()
        logging.getLogger("adapters").setLevel(logging.ERROR)
        tokenizer = transformers.AutoTokenizer.from_pretrained("allenai/specter2_base")
        model = AutoAdapterModel.from_pretrained("allenai/specter2_base")
        model.load_adapter("allenai/specter2", source="hf", load_as="proximity", set_active=True)
        model.eval()
    except Exception as err:  # library missing, no internet for the first download, version clash
        print(f"Local model unavailable ({type(err).__name__}: {err}). Using word matching only.")
        return None

    def embed(pairs):
        texts = [title + tokenizer.sep_token + text for title, text in pairs]
        rows = []
        with torch.no_grad():
            for i in range(0, len(texts), 8):
                batch = tokenizer(texts[i:i + 8], padding=True, truncation=True, max_length=512,
                                  return_tensors="pt", return_token_type_ids=False)
                rows.append(torch.nn.functional.normalize(model(**batch).last_hidden_state[:, 0, :], dim=1))
        return torch.cat(rows).numpy()

    return embed


def semantic_similarity(embed, sections, docs):
    """Meaning-based similarity of (title, text) docs to sections, shape (docs, sections), 0..1."""
    chunks, owner = [], []
    for j, (heading, body) in enumerate(sections):
        words = body.split()
        for i in range(0, max(len(words), 1), CHUNK_STEP):
            chunks.append((heading, " ".join(words[i:i + CHUNK_WORDS])))
            owner.append(j)
    raw = embed(docs) @ embed(chunks).T
    owner = np.array(owner)
    best_chunk = np.stack([raw[:, owner == j].max(axis=1) for j in range(len(sections))], axis=1)
    return np.clip((best_chunk - SEMANTIC_FLOOR) / max(best_chunk.max() - SEMANTIC_FLOOR, 1e-9), 0, 1)


def level(fraction):
    return "strong" if fraction >= GOOD else "possible" if fraction >= MEDIUM else "weak"


def category(fraction, has_abstract):
    # never call a paper a weak match on its title alone
    return "check" if level(fraction) == "weak" and not has_abstract else level(fraction)


def pdf_path(paper, folder):
    """First PDF of a paper that exists on disk; relative paths are tried next to the library file."""
    for raw in (paper.get("File Attachments") or "").split(";"):
        path = Path(raw.strip()).expanduser()
        for candidate in (path, Path(folder) / path):
            if candidate.suffix.lower() == ".pdf" and candidate.is_file():
                return candidate.resolve()
    return None


def pdf_text(pdf):
    return run_tool([tool("pdftotext"), "-enc", "UTF-8", pdf, "-"])


def thumbnail(pdf, page_no, out_stem):
    """Render one PDF page to <out_stem>.png."""
    run_tool([tool("pdftoppm"), "-f", page_no, "-l", page_no, "-r", 70, "-png", "-singlefile", pdf, out_stem])


def page_blocks(pdf, page_no):
    """Return (page height, [(x0, y0, x1, y1, text)]) for one PDF page, in points."""
    out = run_tool([tool("pdftotext"), "-enc", "UTF-8", "-bbox-layout", "-f", page_no, "-l", page_no, pdf, "-"])
    size = re.search(r'<page width="[\d.]+" height="([\d.]+)"', out)
    blocks = [(*map(float, m.groups()[:4]),
               html.unescape(" ".join(re.findall(r"<word[^>]*>(.*?)</word>", m.group(5)))))
              for m in BLOCK.finditer(out)]
    return (float(size.group(1)) if size else 0.0), blocks


def figure_box(page_height, blocks, fig_no):
    """Return (left, top, right, bottom) around a figure and its caption, or None if unsure."""
    # ponytail: assumes the caption sits under its figure and running text sits above it.
    # Any other layout returns None and the report shows the whole page instead; move to
    # pdffigures2 or GROBID if too many figures fall back.
    caption = next((b for b in blocks if (m := re.match(FIG, b[4])) and m.group(2) == fig_no), None)
    upright = [b for b in blocks if b[3] - b[1] <= 2 * (b[2] - b[0])]  # drop vertical margin text
    if not caption or caption not in upright:
        return None
    x0, y0, x1, y1, _ = caption
    body = [b for b in upright if b is not caption and len(b[4].split()) >= BODY_WORDS]
    left, right = min(b[0] for b in upright), max(b[2] for b in upright)
    header = max((b[3] for b in upright if b[3] < 0.08 * page_height), default=0.03 * page_height)

    def top_of(lo, hi):  # bottom of the nearest running text above the caption, else the page header
        return max((b[3] for b in body if b[3] <= y0 + 2 and b[2] > lo + 5 and b[0] < hi - 5), default=header)

    top = top_of(left, right)
    if x1 - x0 < 0.55 * (right - left):  # caption is one column wide
        mid, in_left = (left + right) / 2, (x0 + x1) / 2 < (left + right) / 2
        own = (left, mid) if in_left else (mid, right)
        top = top_of(*own)
        beside = sum(max(0, min(b[3], y0) - max(b[1], top)) for b in body
                     if ((b[0] + b[2]) / 2 < mid) != in_left)
        if beside > 0.3 * (y0 - top):  # running text next to the figure: it lives in one column
            left, right = own
        # otherwise the figure spans both columns and only its caption is narrow
    if not 40 <= y0 - top <= 0.85 * page_height:
        return None  # nothing above the caption (caption-on-top layout) or implausibly tall
    return left, top, right, y1


def crop_figure(pdf, page_no, fig_no, out_stem):
    """Render just the figure to <out_stem>.png; return False if it could not be isolated."""
    box = figure_box(*page_blocks(pdf, page_no), fig_no)
    if not box:
        return False
    scale, pad = 150 / 72, 6
    left, top, right, bottom = box
    run_tool([tool("pdftoppm"), "-f", page_no, "-l", page_no, "-r", 150,
              "-x", int(max(left - pad, 0) * scale), "-y", int(max(top + 2, 0) * scale),
              "-W", int((right - left + 2 * pad) * scale), "-H", int((bottom - top + pad) * scale),
              "-png", "-singlefile", pdf, out_stem])
    return Path(f"{out_stem}.png").is_file()


def read_opening(text):
    """From a paper's PDF text return (text before the introduction, introduction, how it was found)."""
    # ponytail: finds an "Introduction" heading and reads to the next heading. Papers without that
    # heading (letters, some journals) fall back to their opening pages. GROBID would split sections
    # properly if this proves too rough.
    match = INTRO_HEAD.search(text)
    if match:
        rest = text[match.end():]
        end = INTRO_END.search(rest)
        words = rest[:end.start() if end else len(rest)].split()
        if len(words) >= 50:
            return " ".join(text[:match.start()].split()[-400:]), " ".join(words[:INTRO_WORDS]), "introduction"
    words = text.split()[:INTRO_WORDS]
    return "", " ".join(words), "opening pages" if words else ""


def find_captions(text):
    """Return [(fig_no, page_no, caption)] from pdftotext output (pages split by \\f)."""
    # ponytail: regex on text blocks starting "Fig. N". Misses scanned PDFs and
    # odd layouts; upgrade to GROBID or pdffigures2 if too many figures are missed.
    found = {}
    for pattern in CAPTION_PATTERNS:
        for page_no, page in enumerate(text.split("\f"), 1):
            for m in pattern.finditer(page):
                block = page[m.start(1):m.start(1) + 800].split("\n\n")[0]
                found.setdefault(m.group(2), (m.group(2), page_no, " ".join(block.split())))
    return sorted(found.values(), key=lambda c: c[1])


def openalex(doi):
    """Return (corresponding author names, licence) for a DOI; ([], None) if unknown."""
    doi = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", (doi or "").strip())
    if not doi:
        return [], None
    url = ("https://api.openalex.org/works/doi:" + urllib.parse.quote(doi)
           + "?select=authorships,primary_location")
    try:
        with urllib.request.urlopen(url, timeout=15) as response:
            work = json.load(response)
    except (OSError, ValueError):  # offline, 404, rate limit, bad JSON
        return [], None
    names = [(a.get("author") or {}).get("display_name")
             for a in work.get("authorships") or [] if a.get("is_corresponding")]
    return [n for n in names if n], (work.get("primary_location") or {}).get("license")


def match_score(name, email):
    """How well an email's local part fits an author name: 0 none, 1-2 weak, 3+ strong."""
    local = re.sub(r"[^a-z]", "", email.split("@")[0].lower())
    tokens = [unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode().lower()
              for t in re.split(r"[\W\d_]+", name)]
    tokens = [t for t in tokens if t]
    if not tokens or not local:
        return 0
    surname, given = tokens[-1], tokens[:-1]
    score = 3 * (len(surname) >= 3 and surname in local)
    score += 2 * any(len(g) >= 3 and g in local for g in given)
    if given:  # initials, given-first (hmpark) or surname-first (gwl)
        score += local[0] == given[0][0] or local.startswith(surname[0] + given[0][0])
    return score


def find_emails(text, names):
    """Return (emails, how they were chosen) for a paper's corresponding author(s)."""
    # ponytail: name/marker heuristics on PDF text. Wrong when the right email is
    # not printed in the PDF; the "how" label says how far to trust the result.
    emails = list(dict.fromkeys(m for m in EMAIL.findall(text)
                                if not m.lower().startswith(NOT_AUTHOR)))
    if not emails:
        return [], "not found"
    best = {n: max(emails, key=lambda mail: match_score(n, mail)) for n in names}
    scores = {n: match_score(n, mail) for n, mail in best.items()}
    named = list(dict.fromkeys(mail for n, mail in best.items() if scores[n]))
    if any(s >= 3 for s in scores.values()):
        return named, "matched to corresponding author name"
    marked = [mail for m in CORRESPONDING.finditer(text)
              for mail in EMAIL.findall(text[m.start():m.start() + 300])]
    marked = [mail for mail in dict.fromkeys(marked) if mail in emails]
    if marked:
        return marked, "printed next to 'corresponding author' in the PDF"
    if named:
        return named, "weak name match, verify"
    return emails, "unsure which author, all emails in the PDF"


def license_note(lic):
    if lic in FREE_LICENSES:
        return "no permission needed, cite the source"
    if lic and lic.startswith("cc-"):
        return "reuse is limited (non-commercial / no-derivatives terms), check before use"
    return ("no open licence found: the publisher probably holds the rights, "
            "so also check the publisher's permissions page")


def citation(paper):
    first, *rest = (paper.get("Author") or "[AUTHORS]").split(";")
    return (f'{first.strip()}{" et al." if rest else ""}, "{paper["Title"]}", '
            f'{paper.get("Publication Title") or "[JOURNAL]"} '
            f'({paper.get("Publication Year") or "[YEAR]"}), DOI: {paper.get("DOI") or "[DOI]"}')


def greeting_name(paper):
    return " and ".join("Dr. " + n for n in paper["corr"]) or "Dr. [AUTHOR NAME]"


def draft_mail(paper, fig_no, body=None):
    """Full draft: fixed To/Subject lines, then the AI-written body or the template."""
    head = (f'To: {", ".join(paper["emails"]) or "[AUTHOR EMAIL]"}\n'
            f'Subject: Permission request to reuse Figure {fig_no} from "{paper["Title"]}"\n\n')
    return head + (body or MAIL_BODY.format(name=greeting_name(paper), fig=fig_no,
                                            citation=citation(paper)))


def read_env_file(path):
    """Read NAME=value lines; blank lines, # comments and a leading "export" are ignored."""
    values = {}
    for line in Path(path).read_text(encoding="utf-8", errors="ignore").splitlines():
        name, sep, value = line.partition("=")
        name = name.strip().removeprefix("export ").strip()
        if sep and name and not name.startswith("#"):
            values[name] = value.strip().strip("'\"")
    return values


def llm_config(url="", model="", key="", file=None):
    """Return (base_url, model, key) if AI mode is configured, else None.

    Each setting comes from the command line, else the environment, else the file.
    """
    saved = read_env_file(file) if file else {}
    base = (url or os.environ.get("LLM_BASE_URL") or saved.get("LLM_BASE_URL", "")).rstrip("/")
    model = model or os.environ.get("LLM_MODEL") or saved.get("LLM_MODEL", "")
    key = key or os.environ.get("LLM_API_KEY") or saved.get("LLM_API_KEY", "")
    if base and model:
        return base, model, key
    if base or model or key:
        print("AI mode off: set both LLM_BASE_URL and LLM_MODEL "
              "(and LLM_API_KEY unless the model is local)")
    return None


def chat(config, system, user):
    """One OpenAI-style chat completion. Returns the reply text, or None on any failure."""
    base, model, key = config
    body = json.dumps({"model": model, "messages": [{"role": "system", "content": system},
                                                    {"role": "user", "content": user}]}).encode()
    headers = {"Content-Type": "application/json"}
    if key:
        headers["Authorization"] = "Bearer " + key
    request = urllib.request.Request(base + "/chat/completions", body, headers)
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            reply = json.load(response)["choices"][0]["message"]["content"]
        if isinstance(reply, str):
            return reply
        print("  AI call failed: reply had no text")
    except urllib.error.HTTPError as err:
        print(f"  AI call failed: HTTP {err.code} {err.read(200).decode(errors='ignore')}")
    except (OSError, ValueError, KeyError, IndexError, TypeError) as err:
        print(f"  AI call failed: {err!r}")
    return None


def parse_list(reply):
    """Pull the JSON array of objects out of a model reply (tolerates fences and chatter)."""
    m = re.search(r"\[.*\]", reply or "", re.S)
    try:
        data = json.loads(m.group(0)) if m else []
    except ValueError:
        return []
    return [d for d in data if isinstance(d, dict)] if isinstance(data, list) else []


def brief(text, words):
    return " ".join((text or "").split()[:words])


def ai_rate_papers(config, sections, papers):
    """Overwrite pct/section/reason on every paper the AI rates; return how many."""
    headings = [h for h, _ in sections]
    draft = "\n\n".join(f"## {h}\n{brief(b, SECTION_WORDS)}" for h, b in sections)
    rated = 0
    for start in range(0, len(papers), BATCH):
        batch = papers[start:start + BATCH]
        listing = json.dumps([{"id": i, "title": p["Title"],
                               "abstract": brief(p.get("abstract") or p.get("Abstract Note"), ABSTRACT_WORDS),
                               "introduction": brief(p.get("intro"), INTRO_AI_WORDS)}
                              for i, p in enumerate(batch)], ensure_ascii=False)
        reply = chat(config, RATE_SYSTEM,
                     f"ARTICLE DRAFT (start of each section):\n{draft}\n\nPAPERS:\n{listing}")
        for item in parse_list(reply):
            i, score = item.get("id"), item.get("score")
            if (type(i) is not int or not 0 <= i < len(batch) or batch[i]["rated_by"] == "ai"
                    or not isinstance(score, (int, float)) or not 0 <= score <= 100):
                continue  # ignore anything malformed; that paper keeps its word-match rating
            p = batch[i]
            p["pct"], p["rated_by"] = score / 100, "ai"
            p["reason"] = str(item.get("reason") or "")[:200]
            if item.get("section") in headings:
                p["section"] = item["section"]
            p["sims"] = [p["pct"] if h == p["section"] else 0.0 for h in headings]
            rated += 1
    return rated


def ai_pick_figures(config, heading, body, figures, candidates, n):
    """Return [(figure index, reason)] chosen from candidates, best first; [] on failure."""
    listing = json.dumps([{"id": i, "paper": figures[i][0]["Title"], "caption": figures[i][3][:400]}
                          for i in candidates], ensure_ascii=False)
    reply = chat(config, FIG_SYSTEM, f"Pick up to {n} figures.\n\nSECTION: {heading}\n"
                                     f"{brief(body, SECTION_WORDS)}\n\nFIGURES:\n{listing}")
    picks = {}
    for item in parse_list(reply):
        if type(item.get("id")) is int and item["id"] in candidates:
            picks.setdefault(item["id"], str(item.get("reason") or "")[:200])
    return list(picks.items())[:n]


def ai_mail_body(config, paper, fig_no, caption, heading, body):
    """Return a personalised email body, or None to fall back to the template."""
    reply = chat(config, MAIL_SYSTEM,
                 f"Address the email to: {greeting_name(paper)}\n"
                 f"Paper: {citation(paper)}\nFigure wanted: Figure {fig_no}\n"
                 f"Its caption: {caption[:600]}\n\n"
                 f"Section of the article where it would appear: {heading}\n{brief(body, 150)}")
    # a reply without the placeholder has probably invented the sender's details
    return reply.strip() if reply and "[YOUR NAME]" in reply else None


def main(draft, library, figs_per_section, config=None, semantic=False):
    lines, references = split_references(read_lines(draft))
    sections = split_sections(lines)
    bodies = [body for _, body in sections]
    words = sum(len(body.split()) for body in bodies)
    if words < MIN_WORDS:
        raise SystemExit(f"The draft has no text to match papers against (found {words} words). "
                         "Check that you picked the right file.")
    papers = read_papers(library)
    if not papers:
        raise SystemExit(f"No rows with a Title column in {library} - is it a Zotero CSV export?")

    have_poppler = bool(tool("pdftotext") and tool("pdftoppm"))
    if not have_poppler:
        print("Poppler not found - papers are judged on their abstract only, and figures and emails are "
              "skipped. " + ("Run setup_windows.bat again." if os.name == "nt" else
                             "Install it, e.g. brew install poppler (macOS) or apt install poppler-utils (Linux)."))
    print(f"Reading the abstract and introduction of {len(papers)} papers...")
    for p in papers:
        p["pdf"] = pdf_path(p, Path(library).parent)
        p["text"] = pdf_text(p["pdf"]) if p["pdf"] and have_poppler else ""
        before_intro, p["intro"], how = read_opening(p["text"])
        # Zotero's abstract when it has one, else whatever the PDF prints before its introduction
        p["abstract"] = p.get("Abstract Note") or before_intro
        p["read"] = " + ".join(filter(None, ["abstract" if p["abstract"] else "", how])) or "title only"

    sim = similarity(bodies, [" ".join([*(p.get(k) or "" for k in PAPER_FIELDS), p["abstract"], p["intro"]])
                              for p in papers])
    # Raw scores shift with draft length and library, so papers are rated
    # relative to the best-matching paper rather than on an absolute number.
    sim, method = sim / (float(sim.max()) or 1.0), "words"
    embed = None
    if semantic:
        print("Loading the local model (the first run downloads about 840 MB)...")
        embed = load_specter()
    if embed:
        # the model reads about 400 words at a time, so abstract and introduction go in separately
        by_abstract = semantic_similarity(embed, sections, [(p["Title"], p["abstract"]) for p in papers])
        by_intro = semantic_similarity(embed, sections, [(p["Title"], brief(p["intro"], CHUNK_WORDS)) for p in papers])
        sim = (sim + np.maximum(by_abstract, by_intro)) / 2
        method = "words+meaning"
    for p, row in zip(papers, sim):
        p["sims"], p["pct"] = row, float(row.max())
        p["section"] = sections[row.argmax()][0]
        p["rated_by"], p["reason"] = method, ""
        p["corr"], p["license"], p["emails"], p["email_how"] = [], None, [], ""
    if config:
        print(f"AI mode on: sending draft excerpts, abstracts and captions to {config[0]} "
              f"(model {config[1]})")
        rated = ai_rate_papers(config, sections, papers)
        print(f"  AI rated {rated}/{len(papers)} papers; the rest keep their word-match rating")
        if not rated:
            print("  AI gave no usable ratings: continuing without AI")
            config = None
    for p in papers:
        p["cat"] = category(p["pct"], p["read"] != "title only")
    papers.sort(key=lambda p: -p["pct"])
    cats = {c: [p for p in papers if p["cat"] == c] for c in LABELS}
    shortlist = cats["strong"] + cats["possible"]
    find_citations(sections, references, papers)
    cited = sum(1 for p in papers if p["cited_in"] or p["in_refs"])

    print(f"Looking up {len(shortlist)} strong/possible matches on OpenAlex and in their PDFs...")
    figures = []  # (paper, fig_no, page_no, caption)
    for p in shortlist:
        p["corr"], p["license"] = openalex(p.get("DOI"))
        p["emails"], p["email_how"] = find_emails(p["text"], p["corr"])
        figures += [(p, *c) for c in find_captions(p["text"])]
    fig_sim = similarity(bodies, [f"{cap} {p['Title']}" for p, _, _, cap in figures]) if figures else []

    out = Path(library).with_name(Path(library).stem + "_report")
    out.mkdir(exist_ok=True)
    with open(out / "papers.csv", "w", newline="", encoding="utf-8-sig") as f:  # the BOM lets Excel show accents
        w = csv.writer(f)
        w.writerow(["rank", "match", "match_pct", "rated_by", "read_before_judging", "reason", "best_section",
                    "cited_in_draft", "title", "year", "doi", "has_abstract", "has_pdf", "corresponding_author",
                    "email", "email_basis", "license"])
        for i, p in enumerate(papers, 1):
            where = "; ".join(sections[j][0] for j in sorted(p["cited_in"])) or ("reference list" if p["in_refs"] else "no")
            w.writerow([i, LABELS[p["cat"]], f"{p['pct']:.0%}", p["rated_by"], p["read"], p["reason"], p["section"],
                        where, p["Title"],
                        p.get("Publication Year", ""), p.get("DOI", ""),
                        "yes" if p.get("Abstract Note") else "no", "yes" if p["pdf"] else "no",
                        "; ".join(p["corr"]), "; ".join(p["emails"]), p["email_how"], p["license"] or ""])

    e = html.escape
    method = (f"an AI model ({e(config[1])})" if config else
              "word overlap plus meaning (SPECTER2 local model)" if embed else "word overlap")
    doc = ["<!doctype html><meta charset='utf-8'><title>CiteFilter suggestions</title><style>"
           "body{font:15px/1.5 system-ui;max-width:1000px;margin:2em auto;padding:0 1em;color:#0f172a}"
           ".fig{display:flex;gap:1em;margin:1em 0;padding:1em;border:1px solid #ccc;border-radius:6px}"
           ".fig img{width:340px;height:auto;max-height:420px;object-fit:contain;border:1px solid #ddd;flex:none}"
           "pre{white-space:pre-wrap;background:#f4f4f4;padding:1em;border-radius:6px}"
           ".warn{color:#a40000}small{color:#666}"
           ".note{background:#eef7f6;border-left:4px solid #1a9e8f;padding:.7em 1em;border-radius:4px}"
           "ul.papers{list-style:none;padding:0}ul.papers li{margin:.45em 0;display:flex;align-items:flex-start}"
           ".pct{flex:none;width:3em;font-weight:600;text-align:right}ul.papers li>span:last-child{flex:1}"
           ".bar{flex:none;width:90px;height:8px;background:#e2e8f0;border-radius:4px;margin:.5em .7em 0}"
           ".bar i{display:block;height:8px;background:#1a9e8f;border-radius:4px}"
           ".tag{font-size:12px;padding:1px 8px;border-radius:10px;white-space:nowrap}"
           ".new{background:#fdebc8;color:#7a4b00}.here{background:#e2e8f0;color:#475569}.else{background:#dbeafe;color:#1e40af}"
           "summary{cursor:pointer;color:#475569}</style>",
           f"<h1>Suggestions for {e(Path(draft).name)}</h1>",
           "<p class='note'><b>Suggestions only. You decide what to cite.</b> Nothing has been removed: "
           "weaker matches are just listed lower.<br>"
           f"<b>Match %</b> is how similar a section is to a paper's title, abstract and introduction, measured by {method}. "
           "It is not a probability that the paper is relevant"
           f"{'' if config or embed else '; 100% is the best-matching paper in your library'}.</p>",
           f"<p>{len(papers)} papers: {len(cats['strong'])} strong match (&ge;{GOOD:.0%}), "
           f"{len(cats['possible'])} possible match (&ge;{MEDIUM:.0%}), {len(cats['weak'])} weak match, "
           f"{len(cats['check'])} to check by hand. {len(figures)} figure captions found.<br>"
           f"Read before judging: abstract and introduction for {sum(1 for p in papers if p['read'] == 'abstract + introduction')} "
           f"papers, abstract and opening pages for {sum(1 for p in papers if 'opening pages' in p['read'])} "
           "(no Introduction heading in the PDF), abstract only for "
           f"{sum(1 for p in papers if p['read'] == 'abstract')} (no PDF), title only for "
           f"{sum(1 for p in papers if p['read'] == 'title only')}.<br>"
           + (f"Citations: {cited} of these papers are already cited in the draft"
              f"{f' (reference list with {len(references)} entries found)' if references else ''}."
              if cited else
              "Citations: none of these papers were found cited in the draft, so all are tagged "
              "&ldquo;not cited yet&rdquo;. Citations written as superscript numbers without a reference "
              "list cannot be detected.") + "</p>"]

    def cite_tag(p, j):
        if j in p["cited_in"]:
            return ("<span class='tag here'>already cited here</span>" if p["cited_in"][j] == "sure" else
                    "<span class='tag here'>probably cited here (same author and year)</span>")
        if p["cited_in"] or p["in_refs"]:
            return "<span class='tag else'>cited elsewhere in your draft</span>"
        return "<span class='tag new'>not cited yet</span>"

    def paper_line(p, j):
        pct = p["sims"][j]
        return (f"<li><span class='pct'>{pct:.0%}</span><span class='bar'><i style='width:{pct:.0%}'></i></span>"
                f"<span>{e(p['Title'])} <small>({e(p.get('Publication Year') or 'n.d.')}) &middot; {LABELS[level(pct)]}"
                f"{'' if p['read'] == 'abstract + introduction' else ' &middot; read: ' + e(p['read'])}</small> "
                f"{cite_tag(p, j)}{'<br><small>' + e(p['reason']) + '</small>' if p['reason'] else ''}</span></li>")

    thumbs, mail_bodies = {}, {}
    for j, (heading, body) in enumerate(sections):
        ranked_papers = sorted(papers, key=lambda p: -p["sims"][j])
        likely = [p for p in ranked_papers if p["sims"][j] >= MEDIUM]
        weaker = [p for p in ranked_papers if p["sims"][j] < MEDIUM]
        doc.append(f"<h2>{e(heading)}</h2><h3>Suggested papers</h3><ul class='papers'>")
        doc += [paper_line(p, j) for p in likely] or ["<li><span><small>No strong or possible matches for this section.</small></span></li>"]
        doc.append(f"</ul><details><summary>{len(weaker)} weaker matches</summary><ul class='papers'>")
        doc += [paper_line(p, j) for p in weaker]
        doc.append("</ul></details>")
        doc.append("<h3>Suggested figures</h3>")
        ranked = [i for i in sorted(range(len(figures)), key=lambda i: -fig_sim[i][j])
                  if fig_sim[i][j] > 0]
        top = [(i, "") for i in ranked[:figs_per_section]]
        if config and ranked:
            top = ai_pick_figures(config, heading, body, figures, ranked[:CANDIDATES],
                                  figs_per_section) or top
        for i, why in top:
            p, fig_no, page_no, caption = figures[i]
            key = (p["pdf"], fig_no)
            if key not in thumbs:
                stem = f"figure{len(thumbs) + 1}"
                cropped = crop_figure(p["pdf"], page_no, fig_no, out / stem)
                if not cropped:
                    thumbnail(p["pdf"], page_no, out / stem)
                thumbs[key] = (stem, cropped)
            stem, cropped = thumbs[key]
            borrowed = ("<p class='warn'>Caption credits another source: this paper may have borrowed the "
                        "figure, so permission must come from the original source named in the caption. "
                        "The draft email below goes to this paper's author, who can point you to it.</p>"
                        if BORROWED.search(caption) else "")
            mail = ""
            if p["license"] not in FREE_LICENSES or borrowed:
                if config and (p["pdf"], fig_no) not in mail_bodies:
                    mail_bodies[p["pdf"], fig_no] = ai_mail_body(config, p, fig_no, caption, heading, body)
                ai_body = mail_bodies.get((p["pdf"], fig_no))
                mail = (f"<details><summary>Draft permission email "
                        f"({'AI-personalised' if ai_body else 'template'})</summary>"
                        f"<pre>{e(draft_mail(p, fig_no, ai_body))}</pre></details>")
            doc.append(f"<div class='fig'><a href='{p['pdf'].as_uri()}#page={page_no}'>"
                       f"<img src='{stem}.png' alt='Figure {e(fig_no)}'></a><div>"
                       f"<b>Fig. {e(fig_no)}, page {page_no}</b> <small>"
                       f"{'<a href=' + chr(39) + stem + '.png' + chr(39) + '>cropped image</a>' if cropped else 'whole page shown: figure could not be isolated'}"
                       f"</small><br>{e(p['Title'])} "
                       f"<small>({e(p.get('Publication Year') or 'n.d.')}, {LABELS[p['cat']]}) {e(p.get('DOI') or '')}</small>"
                       f"<p>{e(caption[:400])}{'&hellip;' if len(caption) > 400 else ''}</p>"
                       f"{'<p><small>Why this figure: ' + e(why) + '</small></p>' if why else ''}"
                       f"<p><small>Corresponding author: {e(', '.join(p['corr']) or 'not identified by OpenAlex')}"
                       f"<br>Email: {e(', '.join(p['emails']) or '-')} ({e(p['email_how'])})"
                       f"<br>Licence: {e(p['license'] or 'unknown')}, {e(license_note(p['license']))}"
                       f"</small></p>{borrowed}{mail}</div></div>")
        if not top:
            doc.append("<p><small>None found.</small></p>")
    for name, label in (("check", "Check by hand: no abstract and no PDF, so judged on the title only"),
                        ("weak", "Weak match in every section")):
        doc.append(f"<h2>{label}</h2><details><summary>{len(cats[name])} papers</summary><ul>")
        doc += [f"<li>{e(p['Title'])} <small>(best match {p['pct']:.0%}, {e(p['section'])})</small></li>"
                for p in cats[name]]
        doc.append("</ul></details>")
    (out / "report.html").write_text("\n".join(doc), encoding="utf-8")

    print(f"{len(papers)} papers vs {len(sections)} sections: {len(cats['strong'])} strong match, "
          f"{len(cats['possible'])} possible, {len(cats['weak'])} weak, {len(cats['check'])} to check; "
          f"{cited} already cited in the draft")
    print(f"{len(figures)} figure captions; emails for {sum(1 for p in shortlist if p['emails'])}"
          f"/{len(shortlist)} strong/possible papers; corresponding author named by OpenAlex for "
          f"{sum(1 for p in shortlist if p['corr'])}")
    print(f"Report: {out / 'report.html'}\nTable:  {out / 'papers.csv'}")
    return out


def demo():
    lines = [(True, "Absorbers"), (False, "microwave absorber reflection loss ferrite composite " * 8),
             (True, "Batteries"), (False, "lithium battery cathode capacity cycling electrolyte " * 8),
             (True, "Stub"), (False, "too short")]
    sections = split_sections(lines)
    assert [h for h, _ in sections] == ["Absorbers", "Batteries"], sections
    sim = similarity([b for _, b in sections],
                     ["Deep learning for image segmentation neural network pixels",
                      "Ferrite composite microwave absorber reflection loss in X band",
                      "Cathode degradation lithium battery capacity fade on cycling"])
    assert sim[0].max() == 0 and sim[1].argmax() == 0 and sim[2].argmax() == 1, sim
    assert [category(f, True) for f in (1.0, 0.6, 0.59, 0.4, 0.39)] == ["strong", "strong", "possible", "possible", "weak"]
    assert category(0.1, False) == "check" and category(0.7, False) == "strong"

    assert looks_like_heading("2. Materials and Methods") and looks_like_heading("Results") and looks_like_heading("My bold title", bold=True)
    assert is_subsection("2.3 Attenuation Mechanisms in 3D Architectures") and not is_subsection("2.3 GHz was used.") \
        and not is_subsection("1. A numbered list item")
    front = [(False, "Title Author Address " * 10), (True, "Intro"), (False, "word " * 40)]
    assert [h for h, _ in split_sections(front)] == ["Intro"]                      # short front matter dropped
    assert split_sections([(False, "word " * 40)])[0][0] == "Whole article"       # no headings at all
    assert not looks_like_heading("We measured the samples twice.") and not looks_like_heading("2 mm thick samples were cut and then polished") \
        and not looks_like_heading("plain short line")
    assert pdf_path({"File Attachments": "/nonexistent/a.pdf; b.html"}, "/tmp") is None

    body, refs = split_references([(True, "Intro"), (False, "text"), (True, "6. References"),
                                   (False, "[1] Kim S. Radar absorbing carbon materials. 2023."), (False, ""),
                                   (False, "[2] Lee J. doi:10.1/ABC"), (True, "Appendix"), (False, "more")])
    assert body == [(True, "Intro"), (False, "text"), (True, "Appendix"), (False, "more")] and len(refs) == 2
    assert split_references([(False, "no list here")]) == ([(False, "no list here")], [])
    assert cited_numbers("as shown [1, 3-5] and [7\u20138]; not [a] or 2020") == {1, 3, 4, 5, 7, 8}
    lib = [{"Title": "Radar absorbing carbon materials", "Author": "Kim, S; Lee, J", "Publication Year": "2023"},
           {"Title": "Something on waveguides and loss", "DOI": "10.1/abc", "Author": "Lee, J", "Publication Year": "2019"},
           {"Title": "Porous carbon foams for shielding", "Author": "Park, H", "Publication Year": "2021"},
           {"Title": "A paper nobody cites in this draft", "Author": "Nobody, A", "Publication Year": "2020"}]
    find_citations([("Intro", "Absorbers matter [1]."), ("Methods", "We follow [2]; see also Park et al. (2021)."),
                    ("Results", "Nothing cited.")], refs, lib)
    assert [p["cited_in"] for p in lib] == [{0: "sure"}, {1: "sure"}, {1: "likely"}, {}], [p["cited_in"] for p in lib]
    assert [p["in_refs"] for p in lib] == ["sure", "sure", "", ""]

    text = ("Intro text.\nFig. 1. The in-text mention continues here\nand on.\n\n"
            "Fig. 1. SEM images of the\nferrite composite.\n\nBody again.\f"
            "Figure 2: Reflection loss curves.\n\nMore body. Fig. 9. Not a caption.\n\n"
            "Fig. 3 | Nature style caption.\n\nFIG. 4. APS style caption.\n\n"
            "(b)\nFig. 5. Caption right under a panel label.\n\nFigure 6. a) lowercase start.")
    caps = find_captions(text)
    assert [(n, pg) for n, pg, _ in caps] == [("1", 1), ("2", 2), ("3", 2), ("4", 2), ("6", 2), ("5", 2)], caps
    assert caps[0][2] == "Fig. 1. SEM images of the ferrite composite.", caps[0]

    words = lambda n: "word " * n
    two_col = [(50, 60, 290, 200, words(40)), (310, 60, 550, 700, words(200)),      # body text, both columns
               (120, 260, 140, 270, "(a)"), (50, 400, 290, 430, "Fig. 3. A caption."),  # label + caption, left
               (50, 440, 290, 700, words(60)), (570, 100, 580, 600, words(8))]         # body below; margin text
    assert figure_box(800, two_col, "3") == (50, 200, 290, 430)                         # stays in its column
    spanning = [b for b in two_col if b[0] != 310] + [(310, 440, 550, 700, words(60))]
    assert figure_box(800, spanning, "3") == (50, 200, 550, 430)                        # widens to both columns
    assert figure_box(800, two_col, "9") is None                                        # no such caption
    on_top = [(50, 60, 550, 380, words(80)), (50, 400, 550, 420, "Fig. 3. Caption above its figure.")]
    assert figure_box(800, on_top, "3") is None                                         # nothing above caption

    paper_text = ("Title of paper\nA. Author\nAbstract\nWe study absorbers.\n\n1. Introduction\n" + "intro words " * 40
                  + "\n\n2. Experimental\nSamples were made.\n")
    before, intro, how = read_opening(paper_text)
    assert how == "introduction" and intro.split() == ["intro", "words"] * 40 and "We study absorbers." in before
    assert read_opening("A letter with no headings.") == ("", "A letter with no headings.", "opening pages")
    assert read_opening("") == ("", "", "")
    assert read_opening("Introduction\ntoo short\n\n2. Methods\nrest of the paper")[2] == "opening pages"

    page = ("A. Coauthor (acoauthor@uni.edu), Hana‐Mi Park, Weilin Gu\n"
            "E-mail: hmpark@uni.example; gwl@inst.example\nContact permissions@publisher.example")
    assert find_emails(page, ["Hana‐Mi Park", "Weilin Gu"]) == (
        ["hmpark@uni.example", "gwl@inst.example"], "matched to corresponding author name")
    assert find_emails(page, []) == (["hmpark@uni.example", "gwl@inst.example"],
                                     "printed next to 'corresponding author' in the PDF")
    assert find_emails("a@b.org and c@d.org", ["Zed Quux"])[1].startswith("unsure")
    assert find_emails("no addresses here", ["X Y"]) == ([], "not found")
    assert "no permission" in license_note("cc-by") and "publisher" in license_note(None)
    paper = {"Title": "T", "Author": "Kim, S; Lee, S", "emails": ["a@b.org"], "corr": ["Jo Bloggs"]}
    assert "[ARTICLE TITLE]" in MAIL_BODY and "[ARTICLE TITLE]" in MAIL_SYSTEM
    mail = draft_mail(paper, "6")
    assert "To: a@b.org" in mail and "Dear Dr. Jo Bloggs" in mail and "Kim, S et al." in mail
    assert draft_mail(paper, "6", "Hello [YOUR NAME]").endswith('from "T"\n\nHello [YOUR NAME]')

    # fake 2-d embeddings: 0 deg = topic A, 90 deg = topic B; "long" needs its 2nd chunk to match
    angle = {"docA": 0, "docB": 90, "docMid": 45, "A": 0, "B": 90}
    def fake_embed(pairs):
        degrees = [angle.get(title, 90 if "late" in text.split()[:CHUNK_STEP] else 0) for title, text in pairs]
        return np.array([[np.cos(np.radians(d)), np.sin(np.radians(d))] for d in degrees])
    floor = SEMANTIC_FLOOR
    globals()["SEMANTIC_FLOOR"] = 0.0
    try:
        sem = semantic_similarity(fake_embed, [("A", "x"), ("B", "y"), ("long", "w " * 260 + "late")],
                                  [("docA", ""), ("docB", ""), ("docMid", "")])
    finally:
        globals()["SEMANTIC_FLOOR"] = floor
    assert sem.shape == (3, 3) and np.allclose(sem[0], [1, 0, 1]) and np.allclose(sem[1], [0, 1, 1], atol=1e-6), sem
    assert np.allclose(sem[2], [0.7071] * 3, atol=1e-3), sem

    with tempfile.NamedTemporaryFile("w", suffix=".env") as tmp:
        tmp.write("# comment\nexport LLM_BASE_URL = http://x/v1/\nLLM_MODEL='m'\nLLM_API_KEY=k=1\n\njunk\n")
        tmp.flush()
        assert read_env_file(tmp.name) == {"LLM_BASE_URL": "http://x/v1/", "LLM_MODEL": "m", "LLM_API_KEY": "k=1"}
        env, os.environ = os.environ, {"LLM_MODEL": "from-env"}
        try:
            assert llm_config(file=tmp.name) == ("http://x/v1", "from-env", "k=1")
            assert llm_config(model="from-cli", file=tmp.name) == ("http://x/v1", "from-cli", "k=1")
            os.environ = {}
            assert llm_config() is None and llm_config(url="http://y", model="m") == ("http://y", "m", "")
        finally:
            os.environ = env

    item = {"data": {"title": "T", "date": "March 2021", "DOI": "10.1/x", "abstractNote": "A",
                     "creators": [{"creatorType": "author", "lastName": "Park", "firstName": "S"},
                                  {"creatorType": "editor", "lastName": "Ed", "firstName": "E"},
                                  {"creatorType": "author", "name": "ACME Lab"}],
                     "tags": [{"tag": "radar"}, {"tag": "carbon"}], "publicationTitle": "J"}}
    row = zotero_row(item, ["/a.pdf", "/b.pdf"])
    assert set(row) == set(ZOTERO_COLUMNS) and row["Author"] == "Park, S; ACME Lab", row
    assert (row["Publication Year"], row["Manual Tags"], row["File Attachments"]) == ("2021", "radar; carbon", "/a.pdf; /b.pdf")
    assert zotero_pdf_path({"key": "K", "data": {"contentType": "text/html"}}) is None

    assert parse_list('Sure!\n```json\n[{"id": 0}, 7]\n```') == [{"id": 0}]
    assert parse_list("no json") == [] and parse_list(None) == [] and parse_list("[oops]") == []
    global chat
    real_chat, secs = chat, [("Absorbers", "x"), ("Batteries", "y")]
    new = lambda: [{"Title": t, "pct": 0.1, "section": "Absorbers", "rated_by": "words", "reason": ""}
                   for t in ("A", "B", "C")]
    try:
        chat = lambda *_: ('[{"id": 1, "score": 90, "section": "Batteries", "reason": "fits"},'
                           '{"id": 2, "score": 900}, {"id": 9, "score": 5}, {"id": 1, "score": 1}]')
        ps = new()
        assert ai_rate_papers(None, secs, ps) == 1
        assert (ps[1]["pct"], ps[1]["section"], ps[1]["sims"], ps[1]["rated_by"]) == (0.9, "Batteries", [0.0, 0.9], "ai")
        assert ps[0]["rated_by"] == ps[2]["rated_by"] == "words" and ps[2]["pct"] == 0.1
        chat = lambda *_: None  # endpoint down: nothing changes
        ps = new()
        assert ai_rate_papers(None, secs, ps) == 0 and all(p["pct"] == 0.1 for p in ps)
        figs = [({"Title": "T"}, "1", 1, "cap")] * 4
        assert ai_pick_figures(None, "h", "b", figs, [0, 1, 2], 2) == []
        chat = lambda *_: '[{"id": 2, "reason": "r"}, {"id": 3}, {"id": 2}, {"id": 0}, {"id": 1}]'
        assert ai_pick_figures(None, "h", "b", figs, [0, 1, 2], 2) == [(2, "r"), (0, "")]
        full = dict(paper, corr=[], emails=[])
        chat = lambda *_: "Dear Dr. Bloggs, I am Jane Smith."  # invented a sender: rejected
        assert ai_mail_body(None, full, "6", "cap", "h", "b") is None
        chat = lambda *_: " Dear Dr. X,\n[YOUR NAME] "
        assert ai_mail_body(None, full, "6", "cap", "h", "b") == "Dear Dr. X,\n[YOUR NAME]"
    finally:
        chat = real_chat
    print("demo ok")


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):  # an old Windows console cannot print every character: never crash on one
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("draft", nargs="?")
    ap.add_argument("library", nargs="?")
    ap.add_argument("--figs", type=int, default=3, help="figures suggested per section (default 3)")
    ap.add_argument("--semantic", action="store_true",
                    help="also match by meaning with the SPECTER2 local model (needs: pip install adapters)")
    ai = ap.add_argument_group("optional AI mode (any OpenAI-style endpoint)")
    ai.add_argument("--llm-url", default="", help="base URL, e.g. http://localhost:11434/v1")
    ai.add_argument("--llm-model", default="", help="model name at that endpoint")
    ai.add_argument("--llm-key", default="", help="API key (stays in your shell history; "
                                                  "prefer the file or LLM_API_KEY)")
    ai.add_argument("--llm-file", help="settings file with LLM_BASE_URL / LLM_MODEL / LLM_API_KEY "
                                       "lines (default: llm.env next to this script)")
    args = ap.parse_args()
    if not args.library:
        demo()
        ap.print_usage()
    else:
        default_file = next((p for p in (APP_DIR / "llm.env", config_dir() / "llm.env") if p.is_file()),
                            APP_DIR / "llm.env")
        if args.llm_file and not Path(args.llm_file).is_file():
            raise SystemExit(f"--llm-file {args.llm_file} not found")
        llm_file = args.llm_file or (default_file if default_file.is_file() else None)
        if args.library.lower().startswith("zotero") and not Path(args.library).is_file():
            try:
                args.library = zotero_library(args.library)
            except ZoteroError as err:
                raise SystemExit(str(err))
        main(args.draft, args.library, args.figs,
             llm_config(args.llm_url, args.llm_model, args.llm_key, llm_file), args.semantic)
