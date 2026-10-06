"""One small real run of CiteFilter, with everything it needs made on the spot.

Used by the automated Windows build to prove the tool works on that system (it runs anywhere):
a Word draft, a two-paper library and a PDF are created in a folder whose name has a space and
accented letters, then the tool runs on them and the report is checked. No internet is used.

    python windows/smoke_test.py
"""
import csv
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import zotero_filter as core  # noqa: E402

INTRO = ("Microwave absorbers made from carbon reduce the reflection of radar waves. "
         "Impedance matching and dielectric loss decide how much of the wave is absorbed. ") * 6


def make_pdf(path, lines):
    """Write a one-page text PDF without any library (enough for pdftotext and pdftoppm)."""
    text = "BT /F1 11 Tf 60 780 Td 14 TL\n" + "".join(
        "(" + line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)") + ") Tj T*\n" for line in lines) + "ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n%s\nendstream" % (len(text), text.encode("latin-1")),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    Path(path).write_bytes(out)


def wrap(text, width=85):
    words, lines = text.split(), [""]
    for word in words:
        if len(lines[-1]) + len(word) + 1 > width:
            lines.append("")
        lines[-1] = (lines[-1] + " " + word).strip()
    return lines


def main():
    import docx

    with tempfile.TemporaryDirectory() as temp:
        folder = Path(temp) / "my libräry é"
        folder.mkdir()
        make_pdf(folder / "paper ü.pdf", [
            "Carbon foam absorbers for radar waves", "Jane Roe and John Doe",
            "Corresponding author. E-mail: jane.roe@uni.example", "",
            "1. Introduction", *wrap(INTRO), "",
            "2. Methods", "Samples were measured in a waveguide between 8 and 12 GHz.", "", "",
            "Fig. 1. Reflection loss of the carbon foam absorber between 8 and 12 GHz."])
        with open(folder / "library.csv", "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=["Title", "Author", "Publication Year", "Abstract Note", "File Attachments"])
            writer.writeheader()
            writer.writerow({"Title": "Carbon foam absorbers for radar waves", "Author": "Roe, Jane; Doe, John",
                             "Publication Year": "2021", "File Attachments": "paper ü.pdf",
                             "Abstract Note": "Carbon foams absorb microwave radiation through dielectric loss."})
            writer.writerow({"Title": "Zebrafish genome editing", "Author": "Poe, Ann", "Publication Year": "2019",
                             "Abstract Note": "Base editors change single nucleotides in zebrafish embryos."})
        draft = docx.Document()
        draft.add_heading("Introduction", 1)
        draft.add_paragraph("Radar absorbing materials based on carbon have been studied widely, for example by "
                            "Roe and Doe (2021). " + INTRO)
        draft.add_heading("Absorber design", 1)
        draft.add_paragraph("The reflection loss of a carbon foam absorber depends on its thickness. " + INTRO)
        draft.save(str(folder / "draft é.docx"))

        out = core.main(str(folder / "draft é.docx"), str(folder / "library.csv"), 2)
        rows = list(csv.DictReader(open(out / "papers.csv", encoding="utf-8-sig")))
        report = (out / "report.html").read_text(encoding="utf-8")
        checks = {
            "report and table written": (out / "report.html").is_file() and len(rows) == 2,
            "right paper ranked first as a strong match": rows[0]["title"].startswith("Carbon foam") and rows[0]["match"] == "strong match",
            "PDF in an accented folder was read (abstract + introduction)": rows[0]["read_before_judging"] == "abstract + introduction",
            "email read from the PDF": rows[0]["email"] == "jane.roe@uni.example",
            "citation in the draft detected": rows[0]["cited_in_draft"] == "Introduction",
            "figure caption found": "Fig. 1" in report and "Reflection loss of the carbon foam" in report,
            "figure image written": (out / "figure1.png").is_file() and (out / "figure1.png").stat().st_size > 500,
            "unrelated paper not suggested as strong": rows[1]["match"] != "strong match",
        }
        for name, ok in checks.items():
            print(("PASS  " if ok else "FAIL  ") + name)
        if not all(checks.values()):
            print("\nfirst row:", rows[0])
            raise SystemExit(1)
    print("smoke test ok")


if __name__ == "__main__":
    main()
